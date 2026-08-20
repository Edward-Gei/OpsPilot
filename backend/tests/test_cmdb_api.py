"""M2 CMDB 接口测试：主机/应用 CRUD、权限矩阵、IP 冲突、删除保护、suggest、Excel。"""
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import update

from app.models.cmdb import Application, Host
from tests.conftest import auth_header, login_for_tokens


HOST_PAYLOAD = {
    "hostname": "web-01",
    "ip": "10.0.0.1",
    "public_ip": "203.0.113.1",
    "project": "tradingkey",
    "ri": "ri-cn-001",
    "host_series": "C7",
    "platform": "阿里云",
    "region": "华东1",
    "os": "CentOS 7.9",
    "cpu_cores": 4,
    "memory_gb": 8,
    "disk_gb": 100,
    "environment": "prod",
    "status": "online",
    "ssh_port": 22,
    "description": "测试主机",
}


async def _create_host(client, headers, **overrides) -> int:
    """测试辅助：建主机返回 id。"""
    payload = {**HOST_PAYLOAD, **overrides}
    resp = await client.post("/api/v1/cmdb/hosts", json=payload, headers=headers)
    body = resp.json()
    assert body["code"] == 0, f"建主机失败: {body}"
    return body["data"]["id"]


def _make_xlsx(rows: list[list]) -> bytes:
    """测试辅助：构造导入用 xlsx（首行表头随意，导入从第 2 行读）。"""
    wb = Workbook()
    ws = wb.active
    ws.append([
        "主机名", "内网IP地址", "公网IP地址", "项目", "RI", "主机系列", "平台", "区域",
        "操作系统", "CPU", "内存", "磁盘", "环境", "状态", "端口", "说明",
    ])
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestHostCrud:
    async def test_host_list_sorts_by_created_at(self, client, db_factory):
        """主机列表按创建时间正反序排序。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        first_id = await _create_host(client, headers, ip="10.0.1.1")
        second_id = await _create_host(client, headers, ip="10.0.1.2")
        created_at = datetime(2026, 1, 1, 8, 0, 0)
        async with db_factory() as session:
            await session.execute(
                update(Host).where(Host.id == first_id).values(created_at=created_at)
            )
            await session.execute(
                update(Host)
                .where(Host.id == second_id)
                .values(created_at=created_at + timedelta(days=1))
            )
            await session.commit()

        resp = await client.get(
            "/api/v1/cmdb/hosts", params={"sort_by": "created_at", "sort_order": "asc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [first_id, second_id]

        resp = await client.get(
            "/api/v1/cmdb/hosts", params={"sort_by": "created_at", "sort_order": "desc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [second_id, first_id]

    async def test_host_crud_and_ip_conflict(self, client):
        """主机增查改删全链路 + IP 冲突 40901。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        host_id = await _create_host(client, headers)

        # 重复 IP -> 40901
        resp = await client.post("/api/v1/cmdb/hosts", json=HOST_PAYLOAD, headers=headers)
        assert resp.json()["code"] == 40901

        # 列表筛选：keyword 命中 IP
        resp = await client.get("/api/v1/cmdb/hosts", params={"keyword": "10.0.0"}, headers=headers)
        body = resp.json()
        assert body["data"]["total"] == 1
        assert body["data"]["items"][0]["hostname"] == "web-01"

        # 编辑：改状态与 IP
        resp = await client.put(
            f"/api/v1/cmdb/hosts/{host_id}",
            json={**HOST_PAYLOAD, "ip": "10.0.0.2", "status": "maintenance"},
            headers=headers,
        )
        assert resp.json()["code"] == 0
        resp = await client.get(f"/api/v1/cmdb/hosts/{host_id}", headers=headers)
        detail = resp.json()["data"]
        assert detail["ip"] == "10.0.0.2"
        assert detail["status"] == "maintenance"
        # 资源配置字段回显
        assert detail["os"] == "CentOS 7.9"
        assert (detail["cpu_cores"], detail["memory_gb"], detail["disk_gb"]) == (4, 8, 100)
        assert detail["apps"] == []

        # 删除属于独立高风险权限，由 admin 执行
        admin_headers = auth_header(await login_for_tokens(client, "admin"))
        resp = await client.delete(f"/api/v1/cmdb/hosts/{host_id}", headers=admin_headers)
        assert resp.json()["code"] == 0
        resp = await client.get(f"/api/v1/cmdb/hosts/{host_id}", headers=headers)
        assert resp.json()["code"] == 40401

    async def test_host_invalid_ip(self, client):
        """IP 格式不合法 -> 40001。"""
        tokens = await login_for_tokens(client, "ops1")
        resp = await client.post(
            "/api/v1/cmdb/hosts",
            json={**HOST_PAYLOAD, "ip": "999.1.1.1"},
            headers=auth_header(tokens),
        )
        assert resp.json()["code"] == 40001

    async def test_host_extended_fields_and_validation(self, client):
        """项目、RI、主机系列和公网 IP 可维护，项目默认值与校验生效。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        host_id = await _create_host(client, headers)
        detail = (await client.get(f"/api/v1/cmdb/hosts/{host_id}", headers=headers)).json()["data"]
        assert (detail["project"], detail["ri"], detail["host_series"], detail["public_ip"]) == (
            "tradingkey", "ri-cn-001", "C7", "203.0.113.1",
        )

        edited = {
            **HOST_PAYLOAD,
            "project": "mitrade",
            "ri": "ri-cn-002",
            "host_series": "C8",
            "public_ip": "2001:db8::1",
        }
        assert (await client.put(f"/api/v1/cmdb/hosts/{host_id}", json=edited, headers=headers)).json()["code"] == 0
        detail = (await client.get(f"/api/v1/cmdb/hosts/{host_id}", headers=headers)).json()["data"]
        assert (detail["project"], detail["ri"], detail["host_series"], detail["public_ip"]) == (
            "mitrade", "ri-cn-002", "C8", "2001:db8::1",
        )

        default_payload = {**HOST_PAYLOAD, "ip": "10.0.0.8"}
        default_payload.pop("project")
        default_id = (await client.post("/api/v1/cmdb/hosts", json=default_payload, headers=headers)).json()["data"]["id"]
        default_detail = (await client.get(f"/api/v1/cmdb/hosts/{default_id}", headers=headers)).json()["data"]
        assert default_detail["project"] == "mitrade"

        for payload in (
            {**HOST_PAYLOAD, "ip": "10.0.0.9", "project": "invalid"},
            {**HOST_PAYLOAD, "ip": "10.0.0.10", "public_ip": "not-an-ip"},
        ):
            assert (await client.post("/api/v1/cmdb/hosts", json=payload, headers=headers)).json()["code"] == 40001

    async def test_suggest(self, client):
        """平台和主机系列自动补全去重 + 前缀过滤。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        await _create_host(client, headers, ip="10.1.0.1", platform="阿里云", host_series="C7")
        await _create_host(client, headers, ip="10.1.0.2", platform="腾讯云", host_series="C8")
        await _create_host(client, headers, ip="10.1.0.3", platform="阿里云", host_series="C7")

        resp = await client.get(
            "/api/v1/cmdb/hosts/suggest", params={"field": "platform"}, headers=headers
        )
        assert sorted(resp.json()["data"]["items"]) == ["腾讯云", "阿里云"] or sorted(
            resp.json()["data"]["items"]
        ) == ["阿里云", "腾讯云"]

        resp = await client.get(
            "/api/v1/cmdb/hosts/suggest", params={"field": "platform", "q": "腾"}, headers=headers
        )
        assert resp.json()["data"]["items"] == ["腾讯云"]

        resp = await client.get(
            "/api/v1/cmdb/hosts/suggest", params={"field": "host_series", "q": "C7"}, headers=headers
        )
        assert resp.json()["data"]["items"] == ["C7"]

        # 非法 field -> 40001
        resp = await client.get(
            "/api/v1/cmdb/hosts/suggest", params={"field": "hostname"}, headers=headers
        )
        assert resp.json()["code"] == 40001

    async def test_perm_matrix(self, client, seed):
        """权限矩阵：approver 可读不可写（PRD §2.2 验收抽查项）。"""
        # 给 approver 角色造一个用户：直接用 admin 建
        admin = auth_header(await login_for_tokens(client, "admin"))
        resp = await client.post(
            "/api/v1/users",
            json={
                "username": "appr1", "password": "Passw0rd123!", "display_name": "审批员",
                "role_ids": [seed["roles"]["approver"]],
            },
            headers=admin,
        )
        assert resp.json()["code"] == 0
        # 新建用户首登强制改密：走 40105 流程拿令牌。
        resp = await client.post(
            "/api/v1/auth/login", json={"username": "appr1", "password": "Passw0rd123!"}
        )
        body = resp.json()
        assert body["code"] == 40105
        resp = await client.put(
            "/api/v1/auth/password",
            json={"change_token": body["data"]["change_token"],
                  "old_password": "Passw0rd123!", "new_password": "Approver#2026"},
        )
        assert resp.json()["code"] == 0
        tokens = await login_for_tokens(client, "appr1", "Approver#2026")
        headers = auth_header(tokens)

        # 可读
        resp = await client.get("/api/v1/cmdb/hosts", headers=headers)
        assert resp.json()["code"] == 0
        # 不可写
        resp = await client.post("/api/v1/cmdb/hosts", json=HOST_PAYLOAD, headers=headers)
        assert resp.json()["code"] == 40301


class TestAppCrud:
    async def test_app_list_sorts_by_language(self, client):
        """应用列表按语言正反序排序。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        first = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "Go 服务", "language": "Go", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        second = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "Java 服务", "language": "Java", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        first_id = first.json()["data"]["id"]
        second_id = second.json()["data"]["id"]

        resp = await client.get(
            "/api/v1/cmdb/apps", params={"sort_by": "language", "sort_order": "asc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [first_id, second_id]

        resp = await client.get(
            "/api/v1/cmdb/apps", params={"sort_by": "language", "sort_order": "desc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [second_id, first_id]

    async def test_app_list_sorts_by_created_at(self, client, db_factory):
        """应用列表按创建时间正反序排序。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        first = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "先创建应用", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        second = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "后创建应用", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        first_id = first.json()["data"]["id"]
        second_id = second.json()["data"]["id"]
        created_at = datetime(2026, 1, 1, 8, 0, 0)
        async with db_factory() as session:
            await session.execute(
                update(Application).where(Application.id == first_id).values(created_at=created_at)
            )
            await session.execute(
                update(Application)
                .where(Application.id == second_id)
                .values(created_at=created_at + timedelta(days=1))
            )
            await session.commit()

        resp = await client.get(
            "/api/v1/cmdb/apps", params={"sort_by": "created_at", "sort_order": "asc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [first_id, second_id]

        resp = await client.get(
            "/api/v1/cmdb/apps", params={"sort_by": "created_at", "sort_order": "desc"}, headers=headers
        )
        assert [item["id"] for item in resp.json()["data"]["items"]] == [second_id, first_id]

    async def test_app_crud_with_hosts(self, client):
        """应用 CRUD + 多对多关联 + 名称唯一 + 主机详情反查。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        h1 = await _create_host(client, headers, ip="10.2.0.1", hostname="app-node-1")
        h2 = await _create_host(client, headers, ip="10.2.0.2", hostname="app-node-2")

        resp = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "订单服务", "language": "Java", "deploy_type": "docker",
                  "description": "核心服务", "host_ids": [h1, h2]},
            headers=headers,
        )
        body = resp.json()
        assert body["code"] == 0
        app_id = body["data"]["id"]

        # 名称唯一 -> 40901
        resp = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "订单服务", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        assert resp.json()["code"] == 40901

        # 列表含关联主机数与 IP 清单，不返回资源汇总
        resp = await client.get("/api/v1/cmdb/apps", headers=headers)
        item = resp.json()["data"]["items"][0]
        assert item["host_count"] == 2
        assert not {"cpu_total", "memory_total", "disk_total"} & item.keys()
        assert item["host_ips"] == ["10.2.0.1", "10.2.0.2"]

        # 详情含主机清单；主机详情反查应用
        resp = await client.get(f"/api/v1/cmdb/apps/{app_id}", headers=headers)
        app_detail = resp.json()["data"]
        assert not {"cpu_total", "memory_total", "disk_total"} & app_detail.keys()
        assert {h["id"] for h in app_detail["hosts"]} == {h1, h2}
        resp = await client.get(f"/api/v1/cmdb/hosts/{h1}", headers=headers)
        assert resp.json()["data"]["apps"][0]["name"] == "订单服务"

        # 编辑：全量替换关联为仅 h2
        resp = await client.put(
            f"/api/v1/cmdb/apps/{app_id}",
            json={"name": "订单服务v2", "language": "Go", "deploy_type": "k8s", "host_ids": [h2]},
            headers=headers,
        )
        assert resp.json()["code"] == 0
        resp = await client.get(f"/api/v1/cmdb/apps/{app_id}", headers=headers)
        detail = resp.json()["data"]
        assert detail["name"] == "订单服务v2"
        assert [h["id"] for h in detail["hosts"]] == [h2]

        # 不存在的主机 id -> 40001
        resp = await client.put(
            f"/api/v1/cmdb/apps/{app_id}",
            json={"name": "订单服务v2", "deploy_type": "k8s", "host_ids": [99999]},
            headers=headers,
        )
        assert resp.json()["code"] == 40001

    async def test_app_project_type_default_update_and_filter(self, client):
        """项目类型默认前端，支持更新和列表筛选，并拒绝枚举外的值。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        frontend = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "前端门户", "deploy_type": "docker", "project_type": "frontend", "host_ids": []},
            headers=headers,
        )
        backend = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "订单后端", "deploy_type": "shell", "project_type": "backend", "host_ids": []},
            headers=headers,
        )
        defaulted = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "默认应用", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )
        frontend_id = frontend.json()["data"]["id"]
        backend_id = backend.json()["data"]["id"]
        defaulted_id = defaulted.json()["data"]["id"]

        detail = await client.get(f"/api/v1/cmdb/apps/{defaulted_id}", headers=headers)
        assert detail.json()["data"]["project_type"] == "frontend"

        resp = await client.put(
            f"/api/v1/cmdb/apps/{backend_id}",
            json={"name": "订单后端", "deploy_type": "shell", "project_type": "frontend", "host_ids": []},
            headers=headers,
        )
        assert resp.json()["code"] == 0
        resp = await client.get(
            "/api/v1/cmdb/apps", params={"project_type": "frontend"}, headers=headers
        )
        assert {item["id"] for item in resp.json()["data"]["items"]} == {
            frontend_id, backend_id, defaulted_id,
        }

        resp = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "非法类型", "deploy_type": "shell", "project_type": "mobile", "host_ids": []},
            headers=headers,
        )
        assert resp.json()["code"] == 40001

    async def test_app_only_prod_hosts(self, client):
        """应用仅可关联生产环境主机：非生产主机 -> 40001。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        dev_host = await _create_host(client, headers, ip="10.6.0.1", environment="demo")
        resp = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "库存服务", "deploy_type": "docker", "host_ids": [dev_host]},
            headers=headers,
        )
        body = resp.json()
        assert body["code"] == 40001
        assert "生产环境" in body["message"]

    async def test_delete_protection(self, client):
        """删除保护：主机被应用关联时 42201，解除后可删；应用删除后关联清理。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        host_id = await _create_host(client, headers, ip="10.3.0.1")
        resp = await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "支付服务", "deploy_type": "shell", "host_ids": [host_id]},
            headers=headers,
        )
        app_id = resp.json()["data"]["id"]

        admin_headers = auth_header(await login_for_tokens(client, "admin"))
        # 主机被应用引用 -> 42201 且 message 指明引用方
        resp = await client.delete(f"/api/v1/cmdb/hosts/{host_id}", headers=admin_headers)
        body = resp.json()
        assert body["code"] == 42201
        assert "支付服务" in body["message"]

        # 删应用后主机可删
        resp = await client.delete(f"/api/v1/cmdb/apps/{app_id}", headers=admin_headers)
        assert resp.json()["code"] == 0
        resp = await client.delete(f"/api/v1/cmdb/hosts/{host_id}", headers=admin_headers)
        assert resp.json()["code"] == 0


class TestHostExcel:
    async def test_template_download(self, client):
        """模板下载：xlsx 内容类型 + 表头正确。"""
        tokens = await login_for_tokens(client, "ops1")
        resp = await client.get(
            "/api/v1/cmdb/hosts/import-template", headers=auth_header(tokens)
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        ws = load_workbook(BytesIO(resp.content)).active
        assert ws.cell(row=1, column=1).value == "主机名*"
        assert [ws.cell(row=1, column=idx).value for idx in range(1, 7)] == [
            "主机名*", "内网IP地址*", "公网IP地址", "项目", "RI", "主机系列",
        ]

    async def test_import_and_failed_rows(self, client):
        """导入：成功行入库、失败行返回明细（行号+原因）。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        await _create_host(client, headers, ip="10.4.0.9", hostname="exists")

        content = _make_xlsx([
            ["imp-1", "10.4.0.1", "203.0.113.11", "tradingkey", "ri-1", "C7", "华为云", "华北", "Ubuntu 22.04", 8, 16, 200, "prod", "online", 22, ""],
            ["imp-default", "10.4.0.8", None, None, None, None, None, None, None, None, None, None, "prod", "online", 22, ""],
            ["imp-2", "bad-ip", None, None, None, None, "华为云", "华北", None, None, None, None, "prod", "online", 22, ""],
            ["imp-3", "10.4.0.3", None, None, None, None, None, None, None, None, None, None, "wrong-env", None, None, ""],
            ["imp-4", "10.4.0.9", None, None, None, None, None, None, None, None, None, None, "demo", None, None, ""],
            ["imp-5", "10.4.0.1", None, None, None, None, None, None, None, None, None, None, "demo", None, None, ""],
            ["imp-6", "10.4.0.6", None, "invalid", None, None, None, None, None, None, None, None, "prod", None, None, ""],
            ["imp-7", "10.4.0.7", "not-an-ip", None, None, None, None, None, None, None, None, None, "prod", None, None, ""],
        ])
        resp = await client.post(
            "/api/v1/cmdb/hosts/import",
            files={"file": ("hosts.xlsx", content,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
        body = resp.json()
        assert body["code"] == 0
        assert body["data"]["success_count"] == 2
        failed = {r["row"]: r["reason"] for r in body["data"]["failed_rows"]}
        assert set(failed) == {4, 5, 6, 7, 8, 9}

        resp = await client.get("/api/v1/cmdb/hosts", params={"keyword": "10.4.0.8"}, headers=headers)
        assert resp.json()["data"]["items"][0]["project"] == "mitrade"

        # upsert 模式：已存在 IP 被更新（含资源配置字段）
        content = _make_xlsx([
            ["exists-new", "10.4.0.9", "203.0.113.19", "mitrade", "ri-updated", "C8", None, None, "Debian 12", 2, 4, 50, "stage", "offline", 2222, ""],
        ])
        resp = await client.post(
            "/api/v1/cmdb/hosts/import?upsert=true",
            files={"file": ("hosts.xlsx", content,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )
        assert resp.json()["data"]["success_count"] == 1
        resp = await client.get("/api/v1/cmdb/hosts", params={"keyword": "10.4.0.9"}, headers=headers)
        item = resp.json()["data"]["items"][0]
        assert item["hostname"] == "exists-new"
        assert item["environment"] == "stage"
        assert item["ssh_port"] == 2222
        assert item["os"] == "Debian 12"
        assert (item["cpu_cores"], item["memory_gb"], item["disk_gb"]) == (2, 4, 50)
        assert (item["public_ip"], item["project"], item["ri"], item["host_series"]) == (
            "203.0.113.19", "mitrade", "ri-updated", "C8",
        )

    async def test_export_with_filter(self, client):
        """导出：按筛选条件输出 xlsx，行数与筛选结果一致。"""
        tokens = await login_for_tokens(client, "ops1")
        headers = auth_header(tokens)
        await _create_host(client, headers, ip="10.5.0.1", environment="prod")
        await _create_host(
            client, headers, ip="10.5.0.2", environment="demo", public_ip="203.0.113.52",
            project="tradingkey", ri="ri-export", host_series="C9",
        )

        resp = await client.get(
            "/api/v1/cmdb/hosts/export", params={"environment": "demo"}, headers=headers
        )
        assert resp.status_code == 200
        ws = load_workbook(BytesIO(resp.content)).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1
        assert rows[0][1] == "10.5.0.2"
        assert rows[0][2:6] == ("203.0.113.52", "tradingkey", "ri-export", "C9")


class TestAppExcel:
    async def test_export_with_project_type_filter(self, client):
        """应用导出按项目类型筛选，并包含项目类型列。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "前端应用", "deploy_type": "docker", "project_type": "frontend", "host_ids": []},
            headers=headers,
        )
        await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "后端应用", "deploy_type": "shell", "project_type": "backend", "host_ids": []},
            headers=headers,
        )

        resp = await client.get(
            "/api/v1/cmdb/apps/export", params={"project_type": "backend"}, headers=headers
        )
        assert resp.status_code == 200
        ws = load_workbook(BytesIO(resp.content)).active
        assert [cell.value for cell in ws[1]][:4] == ["应用名", "开发语言", "部署方式", "项目类型"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1
        assert rows[0][:4] == ("后端应用", None, "shell", "后端")

    async def test_export_with_filter(self, client):
        """应用导出仅包含当前部署方式筛选结果。"""
        headers = auth_header(await login_for_tokens(client, "ops1"))
        await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "Docker 应用", "language": "Go", "deploy_type": "docker", "host_ids": []},
            headers=headers,
        )
        await client.post(
            "/api/v1/cmdb/apps",
            json={"name": "Shell 应用", "language": "Python", "deploy_type": "shell", "host_ids": []},
            headers=headers,
        )

        resp = await client.get(
            "/api/v1/cmdb/apps/export", params={"deploy_type": "docker"}, headers=headers
        )
        assert resp.status_code == 200
        ws = load_workbook(BytesIO(resp.content)).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1
        assert rows[0][0] == "Docker 应用"
