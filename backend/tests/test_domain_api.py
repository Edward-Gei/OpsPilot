"""域名管理 REST 接口的权限与响应边界测试。"""
from dataclasses import replace
from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.constants import DomainProvider
from app.core.security import encrypt_text
from app.dns_providers.base import ProviderUnavailableError, RecordSetDraft, RemoteRecordSet, ZoneRef
from app.models.auth import Permission, Role, RolePermission, User, UserRole
from app.models.domain import DnsZone, DnsZoneBindTask, DnsZoneBindTaskItem
from app.models.job import Credential
from app.schemas.domain import ZoneBindResult
from app.services import domain_bind_task_service, domain_service
from tests.conftest import TEST_PASSWORD_HASH, auth_header, login_for_tokens


REMOTE_A = RemoteRecordSet(
    record_key="api-a",
    record_name="api.example.com",
    record_type="A",
    ttl=300,
    values=("192.0.2.1",),
    read_only_reason=None,
    provider_meta={"locator": {"name": "api.example.com", "type": "A"}},
)
REMOTE_NS = RemoteRecordSet(
    record_key="zone-ns",
    record_name="example.com",
    record_type="NS",
    ttl=300,
    values=("ns1.example.net",),
    read_only_reason="NS 记录不可编辑或删除",
    provider_meta={"locator": {"name": "example.com", "type": "NS"}},
)
REMOTE_TXT = RemoteRecordSet(
    record_key="verify-txt",
    record_name="_verify.example.com",
    record_type="TXT",
    ttl=300,
    values=("private-token-text", "keep-full-value"),
    read_only_reason=None,
    provider_meta={"locator": {"name": "_verify.example.com", "type": "TXT"}},
)


class FakeAdapter:
    """Domain API 测试的本地服务商替身。"""

    def __init__(self):
        self.records = [REMOTE_A, REMOTE_NS]
        self.discovered_zones = [ZoneRef(DomainProvider.AWS_ROUTE53, "Z1", "example.com")]
        self.discover_error: Exception | None = None
        self.list_error: Exception | None = None
        self.discover_calls = 0
        self.list_calls = 0
        self.create_calls: list[RecordSetDraft] = []
        self.replace_calls: list[tuple[RemoteRecordSet, RecordSetDraft]] = []
        self.delete_calls: list[RemoteRecordSet] = []

    async def discover_public_zones(self, credential):
        self.discover_calls += 1
        if self.discover_error:
            raise self.discover_error
        return list(self.discovered_zones)

    async def list_record_sets(self, zone, credential):
        self.list_calls += 1
        if self.list_error:
            raise self.list_error
        return list(self.records)

    async def get_record_set(self, zone, record_key, credential):
        return next((record for record in self.records if record.record_key == record_key), None)

    async def create_simple_record_set(self, zone, record, credential):
        self.create_calls.append(record)
        self.records.append(self._from_draft(record))

    async def replace_simple_record_set(self, zone, before, after, credential):
        self.replace_calls.append((before, after))
        self.records = [
            replace(self._from_draft(after), record_key=before.record_key)
            if record.record_key == before.record_key else record
            for record in self.records
        ]

    async def delete_simple_record_set(self, zone, before, credential):
        self.delete_calls.append(before)
        self.records = [record for record in self.records if record.record_key != before.record_key]

    @staticmethod
    def _from_draft(record: RecordSetDraft) -> RemoteRecordSet:
        return RemoteRecordSet(
            record_key=f"{record.record_name}:{record.record_type}",
            record_name=record.record_name,
            record_type=record.record_type,
            ttl=record.ttl,
            values=tuple(record.values),
            read_only_reason=None,
            provider_meta={"locator": {"name": record.record_name, "type": record.record_type}},
        )


@pytest.fixture
def fake_adapter(monkeypatch):
    adapter = FakeAdapter()
    monkeypatch.setattr(domain_service, "get_adapter", lambda provider: adapter)
    return adapter


async def _make_user_with_permissions(client, db_factory, permissions: set[str]) -> dict:
    async with db_factory() as session:
        rows = await session.execute(select(Permission).where(Permission.code.in_(permissions)))
        permission_ids = {permission.code: permission.id for permission in rows.scalars()}
        role = Role(code="domain-writer", name="域名操作员", description="测试角色")
        session.add(role)
        await session.flush()
        for permission in permissions:
            session.add(RolePermission(role_id=role.id, permission_id=permission_ids[permission]))
        user = User(
            username="domain-writer",
            password_hash=TEST_PASSWORD_HASH,
            display_name="域名操作员",
            source="local",
            status="active",
        )
        session.add(user)
        await session.flush()
        session.add(UserRole(user_id=user.id, role_id=role.id))
        await session.commit()
    return auth_header(await login_for_tokens(client, "domain-writer"))


async def _seed_cloud_credential(
    db_factory,
    *,
    name: str = "route53-prod",
    auth_type: str = "username_password",
) -> Credential:
    async with db_factory() as session:
        credential = Credential(
            name=name,
            login_user="access-key" if auth_type == "username_password" else None,
            auth_type=auth_type,
            secret_enc=encrypt_text("provider-secret"),
            description="Route 53 测试凭据",
        )
        session.add(credential)
        await session.commit()
        return credential


async def test_domain_write_user_can_list_compatible_credentials_without_credential_read(client, db_factory):
    headers = await _make_user_with_permissions(client, db_factory, {"domain:write"})
    await _seed_cloud_credential(db_factory)

    response = await client.get(
        "/api/v1/domains/credentials",
        params={"provider": "aws_route53"},
        headers=headers,
    )

    assert response.status_code == 200
    item = response.json()["data"]["items"][0]
    assert item["auth_type"] == "username_password"
    assert set(item) == {"id", "name", "auth_type", "description"}
    assert (await client.get("/api/v1/credentials", headers=headers)).json()["code"] == 40301


async def test_domain_routes_enforce_separate_read_write_delete_permissions(client):
    headers = auth_header(await login_for_tokens(client, "ops1"))

    assert (await client.get("/api/v1/domains/zones", headers=headers)).json()["code"] == 40301
    assert (
        await client.post(
            "/api/v1/domains/zones/discover",
            json={"provider": "aws_route53", "credential_id": 1},
            headers=headers,
        )
    ).json()["code"] == 40301
    assert (await client.delete("/api/v1/domains/zones/1", headers=headers)).json()["code"] == 40301


async def _bind_zone(client, db_factory, headers, credential_id: int) -> int:
    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential_id,
            "selections": [{"remote_zone_id": "Z1", "description": "生产 Zone"}],
        },
        headers=headers,
    )
    body = response.json()
    assert response.status_code == 202
    assert body["code"] == 0, body
    async with db_factory() as session:
        assert await domain_bind_task_service.process_next_zone_bind_task(session) is True
    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{body['data']['id']}", headers=headers)
    item = detail.json()["data"]["items"][0]
    assert item["status"] == "success"
    return item["zone_id"]


async def test_zone_bind_enqueues_one_hundred_ninety_selections(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    fake_adapter.discovered_zones = [
        ZoneRef(DomainProvider.AWS_ROUTE53, f"Z{index}", f"zone-{index}.example.com")
        for index in range(190)
    ]

    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential.id,
            "selections": [
                {"remote_zone_id": zone.remote_zone_id}
                for zone in fake_adapter.discovered_zones
            ],
        },
        headers=headers,
    )

    body = response.json()
    assert response.status_code == 202
    assert body["code"] == 0, body
    task = body["data"]
    assert task["status"] == "queued"
    assert task["total_count"] == 190
    assert task["success_count"] == task["skipped_count"] == task["failed_count"] == 0
    assert fake_adapter.discover_calls == fake_adapter.list_calls == 0

    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{task['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["data"]["id"] == task["id"]
    assert len(detail.json()["data"]["items"]) == 190
    assert {item["status"] for item in detail.json()["data"]["items"]} == {"pending"}


async def test_zone_bind_task_processes_each_pending_zone(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    fake_adapter.discovered_zones = [
        ZoneRef(DomainProvider.AWS_ROUTE53, "Z1", "one.example.com"),
        ZoneRef(DomainProvider.AWS_ROUTE53, "Z2", "two.example.com"),
    ]

    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential.id,
            "selections": [
                {"remote_zone_id": "Z1", "description": "首批 Zone"},
                {"remote_zone_id": "Z2", "description": "首批 Zone"},
            ],
        },
        headers=headers,
    )
    task_id = response.json()["data"]["id"]

    async with db_factory() as session:
        assert await domain_bind_task_service.process_next_zone_bind_task(session) is True

    assert fake_adapter.discover_calls == 1
    assert fake_adapter.list_calls == 2
    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{task_id}", headers=headers)
    task = detail.json()["data"]
    assert task["status"] == "success"
    assert task["success_count"] == 2
    assert task["skipped_count"] == task["failed_count"] == 0
    assert {item["zone_name"] for item in task["items"]} == {"one.example.com", "two.example.com"}
    assert {item["status"] for item in task["items"]} == {"success"}


async def test_zone_bind_task_reclaims_expired_lease(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    fake_adapter.discovered_zones = [
        ZoneRef(DomainProvider.AWS_ROUTE53, "Z1", "one.example.com"),
        ZoneRef(DomainProvider.AWS_ROUTE53, "Z2", "two.example.com"),
    ]
    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential.id,
            "selections": [{"remote_zone_id": "Z1"}, {"remote_zone_id": "Z2"}],
        },
        headers=headers,
    )
    task_id = response.json()["data"]["id"]

    async with db_factory() as session:
        task = await session.get(DnsZoneBindTask, task_id)
        task.status = "running"
        task.lease_token = "expired-worker"
        task.lease_expires_at = datetime.now() - timedelta(seconds=1)
        first_item = (await session.execute(
            select(DnsZoneBindTaskItem)
            .where(DnsZoneBindTaskItem.task_id == task_id)
            .order_by(DnsZoneBindTaskItem.id)
            .limit(1)
        )).scalar_one()
        first_item.status = "running"
        first_item.lease_token = "expired-worker"
        await session.commit()

    async with db_factory() as session:
        assert await domain_bind_task_service.process_next_zone_bind_task(session) is True

    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{task_id}", headers=headers)
    task = detail.json()["data"]
    assert task["status"] == "success"
    assert task["success_count"] == 2
    assert {item["status"] for item in task["items"]} == {"success"}


async def test_zone_bind_task_keeps_failed_zones_in_ledger_when_discovery_fails(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    fake_adapter.discover_error = ProviderUnavailableError()
    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential.id,
            "selections": [
                {"remote_zone_id": "Z1", "zone_name": "one.example.com"},
                {"remote_zone_id": "Z2", "zone_name": "two.example.com"},
            ],
        },
        headers=headers,
    )
    task_id = response.json()["data"]["id"]

    async with db_factory() as session:
        assert await domain_bind_task_service.process_next_zone_bind_task(session) is True

    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{task_id}", headers=headers)
    task = detail.json()["data"]
    assert task["status"] == "failed"
    assert task["failed_count"] == 2
    assert task["last_error"] == "DNS 服务商暂时不可用，请稍后重试"
    assert {item["status"] for item in task["items"]} == {"failed"}
    assert {item["zone_name"] for item in task["items"]} == {"one.example.com", "two.example.com"}
    assert fake_adapter.list_calls == 0

    zones = await client.get("/api/v1/domains/zones", headers=headers)
    assert zones.status_code == 200
    assert {
        (item["zone_name"], item["record_count"], item["sync_status"], item["last_sync_error"])
        for item in zones.json()["data"]["items"]
    } == {
        ("one.example.com", 0, "failed", "DNS 服务商暂时不可用，请稍后重试"),
        ("two.example.com", 0, "failed", "DNS 服务商暂时不可用，请稍后重试"),
    }


async def test_zone_bind_task_completes_when_item_binding_rolls_back(client, db_factory, fake_adapter, monkeypatch):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)

    async def bind_then_roll_back(session, **kwargs):
        await session.execute(select(DnsZone.id))
        await session.rollback()
        selection = kwargs["selections"][0]
        return [ZoneBindResult(
            remote_zone_id=selection.remote_zone_id,
            status="skipped",
            reason="Zone 已绑定，已跳过",
        )]

    monkeypatch.setattr(domain_bind_task_service.domain_service, "bind_zones", bind_then_roll_back)
    response = await client.post(
        "/api/v1/domains/zones",
        json={
            "provider": "aws_route53",
            "credential_id": credential.id,
            "selections": [{"remote_zone_id": "Z1"}],
        },
        headers=headers,
    )
    task_id = response.json()["data"]["id"]

    async with db_factory() as session:
        assert await domain_bind_task_service.process_next_zone_bind_task(session) is True

    detail = await client.get(f"/api/v1/domains/zone-bind-tasks/{task_id}", headers=headers)
    task = detail.json()["data"]
    assert task["status"] == "success"
    assert task["skipped_count"] == 1
    assert task["items"][0]["status"] == "skipped"


def _make_zone_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zone 导入"
    ws.append(["服务商*", "凭据名称*", "Zone 名称*", "远端 Zone ID*", "描述"])
    for row in rows:
        ws.append(row)
    content = BytesIO()
    wb.save(content)
    return content.getvalue()


async def test_domain_api_discovers_binds_syncs_and_hides_internal_fields(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)

    discovered = await client.post(
        "/api/v1/domains/zones/discover",
        json={"provider": "aws_route53", "credential_id": credential.id},
        headers=headers,
    )
    assert discovered.json()["data"]["items"] == [{"remote_zone_id": "Z1", "zone_name": "example.com"}]

    zone_id = await _bind_zone(client, db_factory, headers, credential.id)
    zones = await client.get("/api/v1/domains/zones", params={"keyword": "生产"}, headers=headers)
    zone = zones.json()["data"]["items"][0]
    assert zone["id"] == zone_id
    assert zone["credential_name"] == "route53-prod"
    assert not {"secret", "secret_enc", "login_user", "provider_meta"}.intersection(zone)

    reads_before_description_update = fake_adapter.list_calls
    updated = await client.put(
        f"/api/v1/domains/zones/{zone_id}",
        json={"description": "核心解析 Zone"},
        headers=headers,
    )
    assert updated.json()["code"] == 0
    assert fake_adapter.list_calls == reads_before_description_update

    fake_adapter.records = [replace(REMOTE_A, values=("192.0.2.2",)), REMOTE_NS]
    synced = await client.post(f"/api/v1/domains/zones/{zone_id}/sync", headers=headers)
    assert synced.json()["code"] == 0
    records = await client.get(
        f"/api/v1/domains/zones/{zone_id}/records",
        params={"record_type": "A", "read_only": "false", "page_size": 1},
        headers=headers,
    )
    record = records.json()["data"]["items"][0]
    assert records.json()["data"]["total"] == 1
    assert record["values"] == ["192.0.2.2"]
    assert "provider_meta" not in record and "record_key" not in record


async def test_record_api_writes_only_simple_current_records(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    zone_id = await _bind_zone(client, db_factory, headers, credential.id)
    records = (await client.get(f"/api/v1/domains/zones/{zone_id}/records", headers=headers)).json()["data"]["items"]
    a_record = next(record for record in records if record["record_type"] == "A")
    ns_record = next(record for record in records if record["record_type"] == "NS")

    read_only = await client.delete(
        f"/api/v1/domains/zones/{zone_id}/records/{ns_record['id']}", headers=headers
    )
    assert read_only.json()["code"] == 42201
    assert fake_adapter.delete_calls == []

    created = await client.post(
        f"/api/v1/domains/zones/{zone_id}/records",
        json={"owner_name": "web", "record_type": "CNAME", "ttl": 300, "values": ["app.example.com"]},
        headers=headers,
    )
    assert created.json()["code"] == 0
    cname = next(
        record for record in (
            await client.get(f"/api/v1/domains/zones/{zone_id}/records", headers=headers)
        ).json()["data"]["items"]
        if record["record_type"] == "CNAME"
    )

    updated = await client.put(
        f"/api/v1/domains/zones/{zone_id}/records/{a_record['id']}",
        json={"ttl": 300, "values": ["192.0.2.3"]},
        headers=headers,
    )
    assert updated.json()["code"] == 0
    assert fake_adapter.replace_calls

    deleted = await client.delete(
        f"/api/v1/domains/zones/{zone_id}/records/{cname['id']}", headers=headers
    )
    assert deleted.json()["code"] == 0
    assert len(fake_adapter.delete_calls) == 1

    fake_adapter.records = [replace(REMOTE_A, values=("192.0.2.99",)), REMOTE_NS]
    drifted = await client.put(
        f"/api/v1/domains/zones/{zone_id}/records/{a_record['id']}",
        json={"ttl": 300, "values": ["192.0.2.4"]},
        headers=headers,
    )
    assert drifted.json()["code"] == 40901


async def test_unbind_removes_only_local_snapshots(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    zone_id = await _bind_zone(client, db_factory, headers, credential.id)

    response = await client.delete(f"/api/v1/domains/zones/{zone_id}", headers=headers)
    assert response.json()["code"] == 0
    assert fake_adapter.delete_calls == []
    async with db_factory() as session:
        assert await session.get(DnsZone, zone_id) is None


async def test_domain_api_rejects_mismatched_credentials_and_sanitizes_provider_failures(
    client,
    db_factory,
    fake_adapter,
):
    headers = auth_header(await login_for_tokens(client, "admin"))
    incompatible = await _seed_cloud_credential(
        db_factory,
        name="google-json",
        auth_type="secret_file",
    )
    mismatch = await client.post(
        "/api/v1/domains/zones/discover",
        json={"provider": "aws_route53", "credential_id": incompatible.id},
        headers=headers,
    )
    assert mismatch.json()["code"] == 42201

    compatible = await _seed_cloud_credential(db_factory)
    fake_adapter.discover_error = RuntimeError("raw provider-secret must not escape")
    failed = await client.post(
        "/api/v1/domains/zones/discover",
        json={"provider": "aws_route53", "credential_id": compatible.id},
        headers=headers,
    )
    assert failed.json()["code"] == 50201
    assert "provider-secret" not in failed.json()["message"]


async def test_zone_import_reports_success_duplicate_and_existing_rows(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    credential = await _seed_cloud_credential(db_factory)
    content = _make_zone_xlsx([
        ["aws_route53", credential.name, "example.com", "Z1", "首次说明"],
        ["aws_route53", credential.name, "example.com", "Z1", "不得覆盖"],
    ])

    imported = await client.post(
        "/api/v1/domains/zones/import",
        files={"file": ("zones.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert imported.status_code == 200
    data = imported.json()["data"]
    assert data["success_count"] == 1
    assert data["skipped_rows"] == [{"row": 3, "reason": "文件内 Zone 重复"}]
    assert data["failed_rows"] == []
    assert fake_adapter.discover_calls == 1

    existing = await client.post(
        "/api/v1/domains/zones/import",
        files={"file": ("zones.xlsx", _make_zone_xlsx([
            ["aws_route53", credential.name, "example.com", "Z1", "不得覆盖"],
        ]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert existing.json()["data"]["skipped_rows"] == [{"row": 2, "reason": "Zone 已绑定，已跳过"}]
    zone = (await client.get("/api/v1/domains/zones", headers=headers)).json()["data"]["items"][0]
    assert zone["description"] == "首次说明"


async def test_zone_template_and_export_keep_full_txt_without_credentials(client, db_factory, fake_adapter):
    headers = auth_header(await login_for_tokens(client, "admin"))
    template = await client.get("/api/v1/domains/zones/import-template", headers=headers)
    assert template.status_code == 200
    template_sheet = load_workbook(BytesIO(template.content)).active
    assert template_sheet.title == "Zone 导入"
    assert [cell.value for cell in template_sheet[1]] == [
        "服务商*", "凭据名称*", "Zone 名称*", "远端 Zone ID*", "描述",
    ]

    credential = await _seed_cloud_credential(db_factory)
    fake_adapter.records.append(REMOTE_TXT)
    await _bind_zone(client, db_factory, headers, credential.id)
    exported = await client.get("/api/v1/domains/zones/export", headers=headers)
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["Zone 台账", "记录集"]
    record_rows = list(workbook["记录集"].iter_rows(min_row=2, values_only=True))
    txt_row = next(row for row in record_rows if row[3] == "TXT")
    assert txt_row[5] == "private-token-text\nkeep-full-value"
    exported_values = "\n".join(
        str(value or "")
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for value in row
    )
    assert "provider-secret" not in exported_values
    assert "access-key" not in exported_values
