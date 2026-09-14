"""Zone 台账 Excel 模板、解析与导出测试。"""
from datetime import datetime
from importlib import import_module
from io import BytesIO

from openpyxl import Workbook, load_workbook


def _make_workbook(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zone 导入"
    sheet.append(["服务商*", "凭据名称*", "Zone 名称*", "远端 Zone ID*", "描述"])
    for row in rows:
        sheet.append(row)
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def test_zone_workbook_template_parse_and_export_shape():
    domain_excel = import_module("app.services.domain_excel")
    template = load_workbook(BytesIO(domain_excel.build_zone_import_template()))
    sheet = template.active
    assert sheet.title == "Zone 导入"
    assert [cell.value for cell in sheet[1]] == [
        "服务商*", "凭据名称*", "Zone 名称*", "远端 Zone ID*", "描述",
    ]

    rows, failures = domain_excel.parse_zone_import_rows(_make_workbook([
        [" aws_route53 ", " route53-prod ", " example.com ", " Z1 ", " 生产 "],
        [None, "route53-prod", "example.com", "Z2", None],
    ]))
    assert [(row.row, row.provider, row.credential_name, row.zone_name, row.remote_zone_id, row.description) for row in rows] == [
        (2, "aws_route53", "route53-prod", "example.com", "Z1", "生产"),
    ]
    assert failures == [{"row": 3, "reason": "服务商不能为空"}]

    exported = domain_excel.export_zones(
        [domain_excel.ZoneExportRow(
            provider="aws_route53",
            credential_name="route53-prod",
            zone_name="example.com",
            remote_zone_id="Z1",
            description="生产",
            record_count=1,
            sync_status="success",
            last_synced_at=datetime(2026, 9, 14, 12, 0, 0),
            last_sync_error=None,
        )],
        [domain_excel.RecordExportRow(
            zone_name="example.com",
            provider="aws_route53",
            record_name="_verify.example.com",
            record_type="TXT",
            ttl=300,
            values=["first", "second"],
            read_only_reason=None,
            last_synced_at=datetime(2026, 9, 14, 12, 0, 0),
        )],
    )
    workbook = load_workbook(BytesIO(exported), read_only=True)
    assert workbook.sheetnames == ["Zone 台账", "记录集"]
    assert list(workbook["记录集"].iter_rows(min_row=2, values_only=True))[0][5] == "first\nsecond"
