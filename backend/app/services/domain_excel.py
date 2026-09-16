"""Zone 台账 Excel 模板、导入编排与双工作表导出。"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.constants import DomainProvider
from app.core.response import BizError, Errors
from app.dns_providers.base import normalize_zone_name
from app.models.domain import DnsRecordSet, DnsZone
from app.models.job import Credential
from app.schemas.domain import ZoneBindSelection
from app.services import domain_service


ZONE_IMPORT_COLUMNS = [
    ("服务商*", "provider", 20),
    ("凭据名称*", "credential_name", 24),
    ("Zone 名称*", "zone_name", 32),
    ("远端 Zone ID*", "remote_zone_id", 32),
    ("描述", "description", 36),
]
MAX_IMPORT_ROWS = 5000


@dataclass(frozen=True)
class ZoneImportRow:
    row: int
    provider: str
    credential_name: str
    zone_name: str
    remote_zone_id: str
    description: str | None


@dataclass(frozen=True)
class ZoneExportRow:
    provider: str
    credential_name: str
    zone_name: str
    remote_zone_id: str
    description: str | None
    record_count: int
    sync_status: str
    last_synced_at: datetime | None
    last_sync_error: str | None


@dataclass(frozen=True)
class RecordExportRow:
    zone_name: str
    provider: str
    record_name: str
    record_type: str
    ttl: int | None
    values: list[str]
    read_only_reason: str | None
    last_synced_at: datetime | None


def build_zone_import_template() -> bytes:
    """生成带固定列与示例行的 Zone 导入模板。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zone 导入"
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(color="FFFFFF", bold=True)
    for index, (title, _, width) in enumerate(ZONE_IMPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = header_fill
        cell.font = header_font
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.append(["aws_route53", "route53-prod", "example.com", "Z1", "示例行，导入前请删除"])
    return _save_workbook(workbook)


def parse_zone_import_rows(content: bytes) -> tuple[list[ZoneImportRow], list[dict]]:
    """按固定列读取并仅清理单元格首尾空白，返回有效行和行级失败。"""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise Errors.param(f"Excel 文件解析失败: {exc.__class__.__name__}") from exc
    try:
        if "Zone 导入" not in workbook.sheetnames:
            raise Errors.param("Excel 必须包含 Zone 导入工作表")
        sheet = workbook["Zone 导入"]
        headers = [_clean_cell(cell) or "" for cell in next(sheet.iter_rows(max_row=1, values_only=True))]
        expected_headers = [title for title, _, _ in ZONE_IMPORT_COLUMNS]
        if headers[:len(expected_headers)] != expected_headers:
            raise Errors.param("Excel 表头不匹配")
        raw_rows = list(sheet.iter_rows(min_row=2, max_col=len(ZONE_IMPORT_COLUMNS), values_only=True))
    finally:
        workbook.close()
    if not raw_rows:
        raise Errors.param("文件无数据行")
    if len(raw_rows) > MAX_IMPORT_ROWS:
        raise Errors.param(f"单次导入最多 {MAX_IMPORT_ROWS} 行，当前 {len(raw_rows)} 行")

    valid_rows: list[ZoneImportRow] = []
    failures: list[dict] = []
    field_names = [field for _, field, _ in ZONE_IMPORT_COLUMNS]
    titles = {field: title.rstrip("*") for title, field, _ in ZONE_IMPORT_COLUMNS}
    for row_number, raw in enumerate(raw_rows, start=2):
        values = {field: _clean_cell(value) for field, value in zip(field_names, raw)}
        if all(value is None for value in values.values()):
            continue
        required = next((field for field in field_names[:4] if not values[field]), None)
        if required:
            failures.append({"row": row_number, "reason": f"{titles[required]}不能为空"})
            continue
        valid_rows.append(ZoneImportRow(
            row=row_number,
            provider=values["provider"],
            credential_name=values["credential_name"],
            zone_name=values["zone_name"],
            remote_zone_id=values["remote_zone_id"],
            description=values["description"],
        ))
    return valid_rows, failures


async def import_zones(
    session: AsyncSession,
    content: bytes,
    *,
    created_by: int,
    audit_context: domain_service.DomainAuditContext,
) -> dict:
    """按凭据分组发现并复用 Zone 绑定逻辑，逐行汇总成功、跳过和失败。"""
    parsed_rows, failed_rows = parse_zone_import_rows(content)
    result = {"success_count": 0, "skipped_rows": [], "failed_rows": list(failed_rows)}
    grouped: dict[tuple[DomainProvider, str], list[tuple[ZoneImportRow, str]]] = {}
    seen_zone_keys: set[tuple[str, str]] = set()
    for row in parsed_rows:
        try:
            provider = DomainProvider(row.provider)
            zone_name = normalize_zone_name(row.zone_name)
        except ValueError:
            result["failed_rows"].append({"row": row.row, "reason": "服务商或 Zone 名称无效"})
            continue
        zone_key = (provider.value, row.remote_zone_id)
        if zone_key in seen_zone_keys:
            result["skipped_rows"].append({"row": row.row, "reason": "文件内 Zone 重复"})
            continue
        seen_zone_keys.add(zone_key)
        grouped.setdefault((provider, row.credential_name), []).append((row, zone_name))

    for (provider, credential_name), rows in grouped.items():
        credential = (await session.execute(
            select(Credential).where(Credential.name == credential_name)
        )).scalar_one_or_none()
        if credential is None:
            _append_group_failures(result, rows, "凭据不存在")
            continue
        try:
            discovered = await domain_service.discover_zones(
                session,
                provider,
                credential.id,
                audit_context,
            )
        except BizError as exc:
            _append_group_failures(result, rows, exc.message)
            continue

        available = {zone.remote_zone_id: zone for zone in discovered}
        selections: list[ZoneBindSelection] = []
        selected_rows: list[ZoneImportRow] = []
        for row, expected_name in rows:
            remote_zone = available.get(row.remote_zone_id)
            if remote_zone is None:
                result["failed_rows"].append({"row": row.row, "reason": "Zone 不可访问或不是公网 Zone"})
            elif remote_zone.zone_name != expected_name:
                result["failed_rows"].append({"row": row.row, "reason": "远端 Zone ID 与 Zone 名称不匹配"})
            else:
                selections.append(ZoneBindSelection(
                    remote_zone_id=row.remote_zone_id,
                    description=row.description,
                ))
                selected_rows.append(row)
        if not selections:
            continue
        try:
            bound = await domain_service.bind_zones(
                session,
                provider=provider,
                credential_id=credential.id,
                selections=selections,
                created_by=created_by,
                audit_context=audit_context,
                discovered=discovered,
            )
        except BizError as exc:
            _append_group_failures(result, [(row, "") for row in selected_rows], exc.message)
            continue
        for row, item in zip(selected_rows, bound):
            if item.status == "success":
                result["success_count"] += 1
            elif item.status == "skipped":
                result["skipped_rows"].append({"row": row.row, "reason": item.reason or "Zone 已跳过"})
            else:
                result["failed_rows"].append({"row": row.row, "reason": item.reason or "Zone 绑定失败"})

    result["skipped_rows"].sort(key=lambda item: item["row"])
    result["failed_rows"].sort(key=lambda item: item["row"])
    return result


async def load_zone_export_rows(
    session: AsyncSession,
    *,
    keyword: str | None = None,
    provider: str | None = None,
    sync_status: str | None = None,
) -> tuple[list[ZoneExportRow], list[RecordExportRow]]:
    """按 Zone 列表同样的筛选条件加载导出行，凭据只取名称。"""
    query = select(DnsZone, Credential.name).outerjoin(Credential, Credential.id == DnsZone.credential_id)
    if keyword:
        like = f"%{keyword}%"
        query = query.where(DnsZone.zone_name.like(like) | DnsZone.description.like(like))
    if provider:
        query = query.where(DnsZone.provider == provider)
    if sync_status:
        query = query.where(DnsZone.sync_status == sync_status)
    zone_rows = (await session.execute(query.order_by(DnsZone.id.desc()))).all()
    zones = [ZoneExportRow(
        provider=zone.provider,
        credential_name=credential_name or "",
        zone_name=zone.zone_name,
        remote_zone_id=zone.remote_zone_id,
        description=zone.description,
        record_count=zone.record_count,
        sync_status=zone.sync_status,
        last_synced_at=zone.last_synced_at,
        last_sync_error=zone.last_sync_error,
    ) for zone, credential_name in zone_rows]
    zone_ids = [zone.id for zone, _ in zone_rows]
    if not zone_ids:
        return zones, []
    record_rows = await session.execute(
        select(DnsRecordSet, DnsZone.zone_name, DnsZone.provider, DnsZone.last_synced_at)
        .join(DnsZone, DnsZone.id == DnsRecordSet.zone_id)
        .where(DnsRecordSet.zone_id.in_(zone_ids))
        .order_by(DnsZone.zone_name, DnsRecordSet.record_name, DnsRecordSet.record_type, DnsRecordSet.id)
    )
    records = [RecordExportRow(
        zone_name=zone_name,
        provider=provider_name,
        record_name=record.record_name,
        record_type=record.record_type,
        ttl=record.ttl,
        values=list(record.values),
        read_only_reason=record.read_only_reason,
        last_synced_at=last_synced_at,
    ) for record, zone_name, provider_name, last_synced_at in record_rows]
    return zones, records


def export_zones(zones: list[ZoneExportRow], records: list[RecordExportRow]) -> bytes:
    """导出 Zone 台账和完整记录集值，任何凭据字段仅保留名称。"""
    workbook = Workbook(write_only=True)
    zone_sheet = workbook.create_sheet("Zone 台账")
    zone_sheet.append([
        "服务商", "凭据名称", "Zone 名称", "远端 Zone ID", "描述", "记录数", "同步状态", "最近同步时间", "错误摘要",
    ])
    for zone in zones:
        zone_sheet.append([
            zone.provider, zone.credential_name, zone.zone_name, zone.remote_zone_id, zone.description,
            zone.record_count, zone.sync_status, _format_datetime(zone.last_synced_at), zone.last_sync_error,
        ])
    record_sheet = workbook.create_sheet("记录集")
    record_sheet.append(["Zone 名称", "服务商", "记录名称", "类型", "TTL", "值", "只读原因", "最近同步时间"])
    for record in records:
        record_sheet.append([
            record.zone_name, record.provider, record.record_name, record.record_type, record.ttl,
            "\n".join(record.values), record.read_only_reason, _format_datetime(record.last_synced_at),
        ])
    return _save_workbook(workbook)


def _clean_cell(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _append_group_failures(result: dict, rows: list[tuple[ZoneImportRow, str]], reason: str) -> None:
    result["failed_rows"].extend({"row": row.row, "reason": reason} for row, _ in rows)


def _format_datetime(value: datetime | None) -> str | None:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else None


def _save_workbook(workbook: Workbook) -> bytes:
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()
