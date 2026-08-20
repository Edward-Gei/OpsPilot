"""主机 Excel 模板 / 导入 / 导出（HOST-06/07）。

设计要点：
- 模板含表头说明行 + 数据校验下拉（环境/状态），降低填写错误率；
- 导入逐行校验（pydantic 复用 HostUpsertRequest），失败行收集 {row, reason} 明细；
- upsert=true 时按 IP 匹配已有主机执行更新，否则重复 IP 记为失败行；
- 导出复用列表筛选语义，openpyxl write_only 流式写，千行级内存友好。
"""
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.response import Errors
from app.models.cmdb import Host
from app.schemas.cmdb import HostUpsertRequest

# 列定义：(表头, 字段名, 列宽)；顺序即模板/导出列序
COLUMNS: list[tuple[str, str, int]] = [
    ("主机名*", "hostname", 22),
    ("内网IP地址*", "ip", 18),
    ("公网IP地址", "public_ip", 18),
    ("项目", "project", 14),
    ("RI", "ri", 16),
    ("主机系列", "host_series", 16),
    ("所属平台", "platform", 16),
    ("所属区域", "region", 16),
    ("操作系统", "os", 16),
    ("CPU核数", "cpu_cores", 10),
    ("内存GB", "memory_gb", 10),
    ("磁盘GB", "disk_gb", 10),
    ("环境*", "environment", 12),
    ("状态", "status", 14),
    ("SSH端口", "ssh_port", 10),
    ("说明", "description", 30),
]
ENV_OPTIONS = "demo,stage,prod"
STATUS_OPTIONS = "online,offline,maintenance"
MAX_IMPORT_ROWS = 5000  # 单次导入上限（验收基准 1000 行 <10s，留足余量）
# Excel 数值单元格可能读为浮点（如 22.0），需去尾清洗的整数字段
INT_FIELDS = ("ssh_port", "cpu_cores", "memory_gb", "disk_gb")


def _column_letter(field: str) -> str:
    """按字段名取模板列号字母（下拉校验定位用，避免列序调整后写死错位）。"""
    idx = next(i for i, (_, f, _) in enumerate(COLUMNS, start=1) if f == field)
    return get_column_letter(idx)


def build_import_template() -> bytes:
    """生成导入模板：蓝色表头 + 环境/状态下拉校验 + 示例行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "主机导入"
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(color="FFFFFF", bold=True)
    for idx, (title, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        ws.column_dimensions[get_column_letter(idx)].width = width
    # 下拉数据校验：环境与状态列按字段名动态定位，覆盖 2~5001 行
    dv_env = DataValidation(type="list", formula1=f'"{ENV_OPTIONS}"', allow_blank=False)
    dv_status = DataValidation(type="list", formula1=f'"{STATUS_OPTIONS}"', allow_blank=True)
    ws.add_data_validation(dv_env)
    ws.add_data_validation(dv_status)
    env_col, status_col = _column_letter("environment"), _column_letter("status")
    dv_env.add(f"{env_col}2:{env_col}{MAX_IMPORT_ROWS + 1}")
    dv_status.add(f"{status_col}2:{status_col}{MAX_IMPORT_ROWS + 1}")
    ws.append([
        "web-server-01", "10.0.0.1", "203.0.113.1", "mitrade", "ri-cn-001", "C7",
        "阿里云", "华东1", "CentOS 7.9", 4, 8, 100, "prod", "online", 22,
        "示例行，导入前请删除",
    ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def import_hosts(
    session: AsyncSession, content: bytes, *, upsert: bool, created_by: int
) -> dict:
    """Excel 导入：逐行校验 + 入库，返回 {success_count, failed_rows}。

    行级失败互不影响；文件内 IP 重复后行报错；upsert=true 时按 IP 更新已有主机。
    """
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 文件格式错误统一转业务异常
        raise Errors.param(f"Excel 文件解析失败: {exc.__class__.__name__}") from exc
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True))
    wb.close()
    if not rows:
        raise Errors.param("文件无数据行")
    if len(rows) > MAX_IMPORT_ROWS:
        raise Errors.param(f"单次导入最多 {MAX_IMPORT_ROWS} 行，当前 {len(rows)} 行")

    # 预取全库 IP -> Host 映射，避免逐行查询（千行级性能保障）
    existing: dict[str, Host] = {
        h.ip: h for h in (await session.execute(select(Host))).scalars()
    }
    seen_ips: set[str] = set()
    success_count = 0
    failed_rows: list[dict] = []
    field_names = [f for _, f, _ in COLUMNS]

    for offset, raw in enumerate(rows):
        row_no = offset + 2  # Excel 实际行号（含表头）
        values = dict(zip(field_names, raw))
        if all(v is None or str(v).strip() == "" for v in values.values()):
            continue  # 跳过整行空白
        try:
            data = _parse_row(values)
        except ValueError as exc:
            failed_rows.append({"row": row_no, "reason": str(exc)})
            continue
        if data["ip"] in seen_ips:
            failed_rows.append({"row": row_no, "reason": f"文件内 IP 重复: {data['ip']}"})
            continue
        seen_ips.add(data["ip"])
        host = existing.get(data["ip"])
        if host is not None:
            if not upsert:
                failed_rows.append({"row": row_no, "reason": f"IP 已存在: {data['ip']}（可勾选\"存在即更新\"）"})
                continue
            for key, value in data.items():
                setattr(host, key, value)
        else:
            host = Host(created_by=created_by, **data)
            session.add(host)
            existing[data["ip"]] = host
        success_count += 1
    await session.flush()
    return {"success_count": success_count, "failed_rows": failed_rows}


def _parse_row(values: dict) -> dict:
    """单行清洗 + pydantic 校验，异常统一转 ValueError（含首条错误原因）。"""
    cleaned = {k: (str(v).strip() if v is not None else None) for k, v in values.items()}
    if not cleaned.get("project"):
        cleaned.pop("project", None)
    for field in ("public_ip", "ri", "host_series"):
        if not cleaned.get(field):
            cleaned[field] = None
    if not cleaned.get("status"):
        cleaned["status"] = "online"
    if not cleaned.get("ssh_port"):
        cleaned["ssh_port"] = "22"
    # 整数字段统一去尾：Excel 数值单元格可能读为 22.0 / 4.0
    for field in INT_FIELDS:
        if cleaned.get(field) and cleaned[field].endswith(".0"):
            cleaned[field] = cleaned[field][:-2]
    try:
        req = HostUpsertRequest(**{k: v for k, v in cleaned.items() if v is not None})
    except ValidationError as exc:
        first = exc.errors()[0]
        field = ".".join(str(x) for x in first.get("loc", []))
        raise ValueError(f"{field}: {first.get('msg', '校验失败')}") from exc
    return req.model_dump()


def export_hosts(hosts: list[Host]) -> bytes:
    """按筛选结果导出 xlsx：write_only 流式模式。"""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("主机列表")
    ws.append([title.rstrip("*") for title, _, _ in COLUMNS] + ["创建时间"])
    for h in hosts:
        ws.append([
            h.hostname, h.ip, h.public_ip, h.project, h.ri, h.host_series, h.platform, h.region, h.os,
            h.cpu_cores, h.memory_gb, h.disk_gb, h.environment,
            h.status, h.ssh_port, h.description,
            h.created_at.strftime("%Y-%m-%d %H:%M:%S") if h.created_at else None,
        ])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
